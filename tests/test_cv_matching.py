from dataclasses import asdict
from io import BytesIO

import pytest

from data_ana.cv_matching import Candidate, candidate_from_text, certificates, extract_cv_text, match_job, minimum_experience, rank_jobs


def job(**overrides):
    return {"source": "itviec", "title": "Python Developer", "company": "Example", "url": "https://example.com/1",
        "requirements": "Required skills: Python, SQL, Docker. At least 2 years of experience building applications.",
        "skills": "Python;SQL;Docker", "description": "Develop backend services", "level": "Unknown", **overrides}


def test_relevant_job_ranks_higher_than_unrelated_job_and_has_explanations():
    candidate = Candidate(skills=["Python", "SQL", "Docker"], years=3)
    results = rank_jobs(candidate, [job(title="Java Developer", skills="Java;Angular", requirements="Java and Angular required. At least 2 years of experience building applications.", url="https://example.com/2"), job()])
    assert results[0]["job"]["url"] == "https://example.com/1"
    assert results[0]["priority"] == "Cao" and results[0]["score"] > results[1]["score"]
    assert results[0]["matched_skills"] == ["Docker", "Python", "SQL"]
    assert results[0]["components"] and "probability" not in results[0]


def test_required_experience_and_english_certificate_gaps_prevent_high_priority():
    role = job(requirements="Python required. At least 5 years of experience. English IELTS 7.0 is required.")
    result = match_job(Candidate(skills=["Python", "SQL", "Docker"], foreign_languages=["Tiếng Anh"], years=1, certificates={"IELTS": 6.0}), role)
    assert result["score"] <= 65 and result["priority"] != "Cao"
    assert result["blockers"] and any("IELTS" in v for v in result["foreign_gaps"])
    unknown = match_job(Candidate(skills=["Python"], foreign_languages=["Tiếng Anh"]), role)
    assert unknown["priority"] != "Cao"


def test_sparse_job_data_cannot_receive_high_rating():
    result = match_job(Candidate(skills=["Python"]), job(requirements="", skills="Python", description=""))
    assert result["score"] <= 60 and result["stars"] <= 3 and result["priority"] != "Cao"
    assert any("Thiếu" in note for note in result["notes"])


def test_numeric_experience_does_not_hide_an_explicit_seniority_gap():
    candidate = Candidate(skills=["Python", "SQL", "Docker"], years=3, level="Junior")
    result = match_job(candidate, job(title="Senior Python Developer"))
    assert result["score"] <= 65 and result["priority"] != "Cao"
    assert any("Senior" in note and "Junior" in note for note in result["blockers"])


def test_no_known_technical_skills_does_not_become_a_high_score_from_english():
    result = match_job(Candidate(foreign_languages=["Tiếng Anh"]), job(title="Employee", skills="", requirements="English is required. Good communication skills are required for this position."))
    assert result["score"] <= 40 and result["priority"] == "Thấp"


def test_cv_parser_preserves_distinctions_and_does_not_sum_education_dates():
    cv = "Junior Developer\nEducation 2018-2022.\nPython3, JavaScript, C++17, Docker.\nEnglish IELTS 6.5.\nNo experience with Java.\n2 years of experience."
    candidate = candidate_from_text(cv)
    assert {"Python", "JavaScript", "C++", "Docker"}.issubset(candidate.skills)
    assert "Java" not in candidate.skills and candidate.years == 2
    assert candidate.certificates == {"IELTS": 6.5}
    assert candidate_from_text("Education 2018-2022. Python development.").years is None


def test_minimum_experience_handles_entry_level_and_ranges():
    assert minimum_experience(job(title="Intern Python Developer", requirements="Python and SQL"))[0] == 0
    assert minimum_experience(job(requirements="Có từ 2-4 năm kinh nghiệm lập trình."))[0] == 2
    assert minimum_experience(job(requirements="Không yêu cầu kinh nghiệm."))[0] == 0


def test_certificates_never_convert_different_frameworks_or_truncate_invalid_scores():
    assert certificates("English B2; IELTS 6.5; TOEIC 850; TOEFL iBT 100") == {"IELTS": 6.5, "TOEIC": 850, "TOEFL iBT": 100, "CEFR": 4.0}
    assert certificates("IELTS 10; TOEIC 1000; TOEFL PBT 600") == {}
    result = match_job(Candidate(skills=["Python", "SQL", "Docker"], years=3, foreign_languages=["Tiếng Anh"], certificates={"TOEIC": 900}), job(requirements="Python, SQL, Docker. English IELTS 6.5 required for this role."))
    assert any("IELTS" in gap and "chưa" in gap for gap in result["foreign_gaps"])
    alternative = match_job(Candidate(skills=["Python", "SQL", "Docker"], years=3, foreign_languages=["Tiếng Anh"], certificates={"TOEIC": 900}), job(requirements="Python SQL Docker. English IELTS 6.5 or TOEIC 800 is required for this role."))
    assert not alternative["foreign_gaps"]


def test_resume_text_and_docx_table_extraction():
    text = "Python developer with SQL and Docker. Two projects in backend development."
    assert extract_cv_text(text.encode(), "resume.txt") == text
    from docx import Document
    document = Document()
    document.add_paragraph("Python developer with two backend projects and teamwork skills.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "SQL Docker"
    stream = BytesIO(); document.save(stream)
    extracted = extract_cv_text(stream.getvalue(), "resume.docx")
    assert "SQL Docker" in extracted and "Python developer" in extracted


def text_pdf():
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=500, height=500)
    font = DictionaryObject({NameObject("/Type"):NameObject("/Font"), NameObject("/Subtype"):NameObject("/Type1"), NameObject("/BaseFont"):NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):font})})
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 20 450 Td (Python developer with SQL, Docker and 2 years of experience.) Tj ET")
    page[NameObject("/Contents")] = contents
    stream = BytesIO(); writer.write(stream)
    return stream.getvalue()


def test_text_pdf_is_read_and_scanned_pdf_is_explained():
    assert "Python developer" in extract_cv_text(text_pdf(), "resume.pdf")
    from pypdf import PdfWriter
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    stream = BytesIO(); writer.write(stream)
    with pytest.raises(ValueError, match="scan"):
        extract_cv_text(stream.getvalue(), "scan.pdf")


@pytest.mark.parametrize("data,filename", [(b"", "empty.txt"), (b"bad PDF", "bad.pdf"), (b"bad DOCX", "bad.docx"), (b"CV file", "resume.exe"), (b"\xff\x01\x02", "bad.txt")])
def test_invalid_uploads_have_controlled_errors(data, filename):
    with pytest.raises(ValueError):
        extract_cv_text(data, filename)


def test_sensitive_personal_fields_are_not_used_for_ranking():
    a = candidate_from_text("Python SQL Docker. 2 years of experience. Female. Age 25. alice@example.com")
    b = candidate_from_text("Python SQL Docker. 2 years of experience. Male. Age 45. bob@example.com")
    assert asdict(a) == asdict(b)
    assert match_job(a, job())["score"] == match_job(b, job())["score"]
