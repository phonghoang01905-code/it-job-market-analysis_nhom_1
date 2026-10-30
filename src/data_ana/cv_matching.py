"""Local, explainable CV-to-job matching. Scores are not hiring probabilities."""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
import zipfile

from data_ana.etl import extract_normalized_skills
from data_ana.job_profile import FOREIGN, SUPPLEMENTARY, extract_profile
from data_ana.languages import extract_languages
from data_ana.utils import infer_level, unaccent

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 100_000
LEVELS = ["Unknown", "Intern", "Fresher", "Junior", "Middle", "Senior"]
WEIGHTS = {"Kỹ năng chuyên môn": 60, "Kinh nghiệm / cấp bậc": 15, "Ngoại ngữ": 15, "Kỹ năng bổ trợ": 10}
CEFR = ["A1", "A2", "B1", "B2", "C1", "C2"]


@dataclass
class Candidate:
    skills: list[str] = field(default_factory=list)
    foreign_languages: list[str] = field(default_factory=list)
    supplementary_skills: list[str] = field(default_factory=list)
    years: float | None = None
    level: str = "Unknown"
    certificates: dict[str, float] = field(default_factory=dict)


def extract_cv_text(data: bytes, filename: str) -> str:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("CV phải có nội dung và không vượt quá 10 MB.")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("Tệp PDF không hợp lệ. Hãy xuất lại CV thành PDF hoặc dán nội dung.")
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF có mật khẩu. Hãy dùng bản không khóa hoặc dán nội dung CV.")
            if len(reader.pages) > 20:
                raise ValueError("Hỗ trợ CV PDF tối đa 20 trang.")
            parts = []
            for page in reader.pages:
                content = page.get_contents()
                if content is not None and len(content.get_data()) > 20 * 1024 * 1024:
                    raise ValueError("Trang PDF quá lớn để đọc. Hãy xuất lại bản nhẹ hơn.")
                parts.append(page.extract_text() or "")
            text = "\n".join(parts)
        elif extension == ".docx":
            from docx import Document
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("DOCX có nội dung giải nén quá lớn. Hãy xuất lại CV nhẹ hơn.")
            document = Document(BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            parts += [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
            text = "\n".join(parts)
        elif extension == ".txt":
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            text = data.decode(encoding)
        else:
            raise ValueError("Chọn CV dạng PDF, DOCX hoặc TXT; hoặc dán nội dung vào ô văn bản.")
    except UnicodeError:
        raise ValueError("Không đọc được mã hóa văn bản. Hãy lưu TXT theo UTF-8 hoặc dán nội dung CV.") from None
    except ValueError:
        raise
    except Exception:
        raise ValueError("Không đọc được tệp CV. Hãy kiểm tra định dạng hoặc dán nội dung văn bản.") from None
    text = text.replace("\x00", "").strip()
    if len(text) < 30 or len(re.findall(r"\w", text)) < 15:
        raise ValueError("Chưa đọc được đủ văn bản. CV có thể là ảnh/scan; hãy dán nội dung CV hoặc dùng PDF có thể chọn chữ.")
    if len(text) > MAX_CHARS:
        raise ValueError("Nội dung CV quá dài; tối đa 100.000 ký tự.")
    return text


def certificates(text: str) -> dict[str, float]:
    result = {}
    for name, pattern, maximum in [
        ("IELTS", r"\bielts\s*(?:overall|score|band|:|-)?\s*(\d(?:[.,]\d)?)(?!\d)", 9),
        ("TOEIC", r"\btoeic\s*(?:score|:|-)?\s*(\d{2,3})(?!\d)", 990),
        ("TOEFL iBT", r"\btoefl\s+ibt\s*(?:score|:|-)?\s*(\d{2,3})(?!\d)", 120),
    ]:
        values = [float(m.replace(",", ".")) for m in re.findall(pattern, text, re.I)]
        values = [v for v in values if 0 < v <= maximum]
        if values:
            result[name] = max(values)
    cefr = re.findall(r"(?:english|tieng anh|cefr).{0,25}?\b([abc][12])\b", unaccent(text))
    cefr += re.findall(r"\b([abc][12])\b.{0,20}?(?:english|cefr)", unaccent(text))
    if cefr:
        result["CEFR"] = float(max(CEFR.index(value.upper())+1 for value in cefr))
    return result


def certificate_label(name: str, value: float) -> str:
    return CEFR[int(value)-1] if name == "CEFR" and 1 <= value <= 6 else f"{value:g}"


def explicit_years(text: str) -> float | None:
    text = unaccent(text)
    # Do not sum dates: they can describe education or overlapping jobs.
    values = re.findall(r"(\d+(?:[.,]\d+)?)\s*\+?\s*(?:nam kinh nghiem|years?(?: of)? (?:work |working |professional )?experience)", text)
    return max((float(v.replace(",", ".")) for v in values if float(v.replace(",", ".")) <= 60), default=None)


def candidate_from_text(text: str) -> Candidate:
    if not text.strip() or len(text) > MAX_CHARS:
        raise ValueError("Nhập nội dung CV, tối đa 100.000 ký tự.")
    # Contact details are not matching features.
    text = re.sub(r"https?://\S+|[\w.+-]+@[\w.-]+", " ", text)
    positive = re.sub(r"(?:no experience with|not familiar with|do not know|don't know|chưa biết|không biết|chưa có kinh nghiệm (?:với|về))\s+[^.;\n]+", " ", text, flags=re.I)
    languages, _ = extract_languages(description=positive)
    tech = sorted(set(extract_normalized_skills("", positive, [])) | set(languages))
    normalized = unaccent(positive)
    foreign = sorted(name for name, pattern in FOREIGN.items() if re.search(pattern, normalized))
    supplementary = sorted(name for name, pattern in SUPPLEMENTARY.items() if re.search(pattern, normalized))
    return Candidate(skills=tech, foreign_languages=foreign, supplementary_skills=supplementary,
        years=explicit_years(positive), level=infer_level(positive[:350]), certificates=certificates(positive))


def minimum_experience(job: dict) -> tuple[float | None, str]:
    text = str(job.get("requirements") or "")
    normalized = unaccent(text)
    match = re.search(r"khong yeu cau kinh nghiem|no (?:prior )?experience required|without experience", normalized)
    if match:
        return 0, text[max(0, match.start()-30):match.end()+70]
    patterns = [
        r"(?:at least|minimum(?: of)?|toi thieu|it nhat|tu)\s+(\d+(?:[.,]\d+)?)\s*\+?\s*(?:years?|nam)",
        r"(\d+(?:[.,]\d+)?)\s*(?:[-–]\s*\d+)?\s*\+?\s*(?:nam kinh nghiem|years?(?: of)? (?:work |working |professional )?experience)",
    ]
    for pattern in patterns:
        for match in re.finditer(pattern, normalized):
            prefix = normalized[max(0, match.start()-35):match.start()]
            if re.search(r"uu tien|preferred|nice.to.have|a plus", prefix):
                continue
            years = float(match.group(1).replace(",", "."))
            if 0 <= years <= 60:
                return years, text[max(0, match.start()-30):match.end()+100]
    level = infer_level(str(job.get("title") or ""))
    return (0, "Vị trí thực tập/fresher") if level in {"Intern", "Fresher"} else (None, "Tin chưa nêu số năm kinh nghiệm rõ ràng")


def _list(value) -> list[str]:
    if isinstance(value, str):
        return [v.strip() for v in value.split(";") if v.strip()]
    return value if isinstance(value, list) else []


def match_job(candidate: Candidate, job: dict) -> dict:
    title = str(job.get("title") or "")
    requirements = str(job.get("requirements") or "")
    body = requirements or str(job.get("description") or "")
    tags = _list(job.get("skills", []))
    languages, _ = extract_languages(title, body, tags)
    tech = set(extract_normalized_skills(title, body, tags)) | set(languages)
    user_tech = set(candidate.skills)
    matched, missing = sorted(tech & user_tech), sorted(tech-user_tech)
    parts, notes, blockers = [], [], []

    def component(name, ratio, detail):
        parts.append({"criterion": name, "weight": WEIGHTS[name], "match": round(ratio*100, 1), "detail": detail})

    if tech:
        component("Kỹ năng chuyên môn", len(matched)/len(tech), f"Khớp {len(matched)}/{len(tech)} công nghệ/kỹ năng được nêu")
    else:
        notes.append("Chưa nhận diện đủ kỹ năng chuyên môn trong tin")
    minimum, experience_excerpt = minimum_experience(job)
    job_level = infer_level(title)
    if job_level == "Unknown":
        job_level = str(job.get("level") or "Unknown")
    if minimum is not None:
        ratio = 1.0 if minimum == 0 else min(candidate.years/minimum, 1) if candidate.years is not None else 0
        component("Kinh nghiệm / cấp bậc", ratio, f"Tin nêu tối thiểu {minimum:g} năm; hồ sơ: {candidate.years:g} năm" if candidate.years is not None else f"Tin nêu {minimum:g} năm; hồ sơ chưa xác nhận số năm")
        if minimum > 0 and (candidate.years is None or candidate.years < minimum):
            blockers.append("Cần đối chiếu/bổ sung kinh nghiệm tối thiểu")
        if job_level in LEVELS[1:] and candidate.level in LEVELS[1:] and LEVELS.index(candidate.level) < LEVELS.index(job_level):
            blockers.append(f"Tin nêu cấp bậc {job_level}, cao hơn cấp bậc hồ sơ {candidate.level}; cần đối chiếu phạm vi công việc")
    elif job_level in LEVELS[1:]:
        if candidate.level in LEVELS[1:]:
            gap = max(0, LEVELS.index(job_level)-LEVELS.index(candidate.level))
            ratio = max(0, 1-gap*0.4)
            component("Kinh nghiệm / cấp bậc", ratio, f"Tin: {job_level}; hồ sơ tự xác nhận: {candidate.level}")
            if gap:
                blockers.append(f"Tin nêu cấp bậc {job_level}, cao hơn cấp bậc hồ sơ")
        else:
            component("Kinh nghiệm / cấp bậc", 0, f"Tin: {job_level}; chưa xác nhận cấp bậc hồ sơ")
            blockers.append("Chưa xác nhận cấp bậc để đối chiếu")
    else:
        notes.append("Tin chưa nêu đủ kinh nghiệm/cấp bậc để chấm")

    profile = extract_profile(title, str(job.get("description") or ""), requirements, tags, str(job.get("benefits") or ""))
    # Use preserved evidence from ETL where original spoken-language tags exist.
    evidence = job.get("_profile_evidence")
    if isinstance(evidence, dict):
        foreign_evidence = evidence.get("foreign_languages", {})
    else:
        foreign_evidence = profile["profile_evidence"]["foreign_languages"]
    foreign = {name: data for name, data in foreign_evidence.items() if data.get("status") != "not_required"}
    foreign_gaps, matched_foreign = [], []
    if foreign:
        points, total = 0.0, 0.0
        for name, data in foreign.items():
            importance = 0.5 if data.get("status") == "preferred" else 1.0
            total += importance
            if name in candidate.foreign_languages:
                ratio = 1.0
                matched_foreign.append(name)
                if name == "Tiếng Anh":
                    required_certs = certificates(data.get("excerpt", ""))
                    alternatives = len(required_certs) > 1 and re.search(r"\bor\b|\bhoac\b", unaccent(data.get("excerpt", "")))
                    alternative_met = alternatives and any(candidate.certificates.get(cert, 0) >= threshold for cert, threshold in required_certs.items())
                    for cert, threshold in required_certs.items():
                        if alternative_met:
                            break
                        value = candidate.certificates.get(cert)
                        if value is None or value < threshold:
                            ratio = min(ratio, 0.5 if value is None else value/threshold)
                            foreign_gaps.append(f"{cert} {certificate_label(cert, threshold)}: hồ sơ chưa có mức/điểm" if value is None else f"{cert} {certificate_label(cert, value)} thấp hơn mức {certificate_label(cert, threshold)} được nêu")
                            if data.get("status") == "required":
                                blockers.append("Chưa đáp ứng/xác nhận chứng chỉ tiếng Anh được yêu cầu")
                elif re.search(r"\btopik\b|\bjlpt\b|\bhsk\b|\bn[1-5]\b", data.get("excerpt", ""), re.I):
                    ratio = 0.5
                    foreign_gaps.append(f"Cần đối chiếu trình độ/chứng chỉ của {name} trong tin gốc")
                    if data.get("status") == "required":
                        blockers.append(f"Chưa tự đối chiếu chứng chỉ {name}")
                points += importance*ratio
            else:
                foreign_gaps.append(f"Chưa thấy {name} trong hồ sơ đã xác nhận")
                if data.get("status") == "required":
                    blockers.append(f"Chưa có thông tin về yêu cầu {name}")
        component("Ngoại ngữ", points/total, "; ".join(foreign_gaps) or "Khớp các ngoại ngữ được nêu; trình độ chưa được kiểm chứng")
    else:
        notes.append("Tin chưa có yêu cầu ngoại ngữ để đối chiếu")

    supporting = set(_list(job.get("supplementary_skills", [])) or profile["supplementary_skills"])
    matched_supp = sorted(supporting & set(candidate.supplementary_skills))
    if supporting:
        component("Kỹ năng bổ trợ", len(matched_supp)/len(supporting), f"Khớp {len(matched_supp)}/{len(supporting)} kỹ năng bổ trợ")
    weight = sum(p["weight"] for p in parts)
    score = sum(p["weight"]*p["match"] for p in parts)/weight if weight else 0
    if not tech:
        score = min(score, 40)
    if len(requirements.strip()) < 40:
        score = min(score, 60)
        notes.append("Thiếu phần yêu cầu chi tiết; điểm tối đa 60/100")
    if blockers:
        score = min(score, 65)
    score = round(score, 1)
    priority = "Cao" if score >= 75 and matched and not blockers else "Vừa" if score >= 45 and matched else "Thấp"
    stars = 5 if score >= 85 else 4 if score >= 70 else 3 if score >= 50 else 2 if score >= 30 else 1 if score > 0 else 0
    return {"score": score, "stars": stars, "priority": priority, "matched_skills": matched, "missing_skills": missing,
        "matched_foreign": matched_foreign, "foreign_gaps": foreign_gaps, "matched_supplementary": matched_supp,
        "experience_min": minimum, "experience_excerpt": experience_excerpt,
        "blockers": sorted(set(blockers)), "notes": notes, "components": parts, "job": job}


def rank_jobs(candidate: Candidate, jobs: list[dict]) -> list[dict]:
    results = [match_job(candidate, job) for job in jobs]
    return sorted(results, key=lambda r: (-r["score"], -len(r["matched_skills"]), str(r["job"].get("source", "")), str(r["job"].get("url", ""))))
